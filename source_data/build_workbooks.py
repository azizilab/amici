"""Assemble the per-panel CSVs written by the figure exporters into one Source Data workbook per figure.

Each workbook starts with a README sheet (provenance of the dataset and checkpoint, and a guide to the
sheets) followed by one sheet per panel. Tables longer than MAX_XLSX_ROWS are written as CSV into a
companion zip (Source_Data_FigN_large_tables.zip) and referenced from the README sheet.

Run from the repository root after the exporters:  python source_data/build_workbooks.py [Fig3 Fig4 ...]
"""

import json
import sys
import zipfile

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import _common as C

MAX_XLSX_ROWS = 150_000
PROVENANCE_FIELDS = [
    ("dataset", "Dataset"),
    ("data", "Preprocessed dataset (repository path)"),
    ("data_md5", "Dataset md5"),
    ("model", "AMICI checkpoint directory (repository path)"),
    ("model_md5", "Checkpoint model.pt md5"),
    ("wandb_sweep", "Weights & Biases sweep"),
    ("wandb_run", "Weights & Biases run"),
    ("config", "Training configuration"),
    ("figshare_model_file_id", "Figshare file id of model.pt (doi:10.6084/m9.figshare.31927956)"),
    ("figshare_data_file_id", "Figshare file id of the dataset"),
    ("gcs_model", "Bucket copy of model.pt"),
    ("gcs_data", "Bucket copy of the dataset"),
]


def write_readme(ws, manifest, large):
    bold = Font(bold=True)
    ws.append([manifest["readme"][0][1]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Provenance"])
    ws.cell(ws.max_row, 1).font = bold
    prov = manifest["provenance"]
    for key, label in PROVENANCE_FIELDS:
        if key in prov:
            ws.append([label, "not deposited" if prov[key] is None else str(prov[key])])
    for key, value in prov.items():
        if key not in dict(PROVENANCE_FIELDS) and key not in ("figures", "labels_key", "seed"):
            ws.append([key, str(value)])
    ws.append([])
    ws.append(["Panel", "Sheet(s)"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for panel, sheets in manifest["readme"][1:]:
        ws.append([panel, sheets])
    ws.append([])
    ws.append(["Sheet", "Rows", "Description"])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for sheet in manifest["sheets"]:
        where = sheet["sheet"]
        if sheet["sheet"] in large:
            where = f"{sheet['sheet']} (CSV in {large[sheet['sheet']]})"
        ws.append([where, sheet["rows"], sheet["description"]])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 120
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build(figure):
    fig_dir = C.PANEL_DIR / figure
    with open(fig_dir / "manifest.json") as handle:
        manifest = json.load(handle)
    C.OUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_path = C.OUT_DIR / f"Source_Data_{figure}.xlsx"
    zip_path = C.OUT_DIR / f"Source_Data_{figure}_large_tables.zip"

    large = {s["sheet"]: zip_path.name for s in manifest["sheets"] if s["rows"] > MAX_XLSX_ROWS}
    if large:
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for sheet in manifest["sheets"]:
                if sheet["sheet"] in large:
                    zf.write(fig_dir / sheet["csv"], arcname=f"Source_Data_{figure}_{sheet['sheet']}.csv")
        print(f"  wrote {zip_path.relative_to(C.REPO)}")
    elif zip_path.exists():
        zip_path.unlink()

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="README", index=False)
        write_readme(writer.book["README"], manifest, large)
        for sheet in manifest["sheets"]:
            if sheet["sheet"] in large:
                continue
            df = pd.read_csv(fig_dir / sheet["csv"], index_col=0 if sheet["index"] else None)
            df.to_excel(writer, sheet_name=sheet["sheet"], index=sheet["index"])
            ws = writer.book[sheet["sheet"]]
            ws.freeze_panes = "B2" if sheet["index"] else "A2"
            for i, col in enumerate(([df.index.name or ""] if sheet["index"] else []) + list(df.columns), start=1):
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, len(str(col)) + 2), 45)
    print(f"  wrote {xlsx_path.relative_to(C.REPO)}")


def main():
    figures = sys.argv[1:] or sorted(p.name for p in C.PANEL_DIR.iterdir() if (p / "manifest.json").exists())
    for figure in figures:
        print(f"Building {figure}")
        build(figure)


if __name__ == "__main__":
    main()
